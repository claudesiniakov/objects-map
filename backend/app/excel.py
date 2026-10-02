"""Чтение входных таблиц (.xlsx/.xls/.csv) и формирование выходных Excel-файлов."""

import csv
import io
import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .config import MAX_ROWS

TEMPLATE_HEADERS = ["ID", "Номер договора", "Кадастровый номер", "Тип", "Название", "Широта", "Долгота", "Адрес",
                    "Радиус, м", "Стоимость, тыс. руб.", "Описание"]


class TableError(ValueError):
    pass


def _clean(v):
    if isinstance(v, str):
        v = v.strip()
        return v if v != "" else None
    return v


def _split_header(rows: list[list]) -> dict:
    """Первая непустая строка — заголовки; пустые строки данных отбрасываются."""
    rows = [[_clean(c) for c in r] for r in rows]
    start = next((i for i, r in enumerate(rows) if any(c is not None for c in r)), None)
    if start is None:
        return {"headers": [], "rows": [], "first_row": 1}
    raw_headers = rows[start]
    width = max(i + 1 for i, c in enumerate(raw_headers) if c is not None)
    headers, seen = [], {}
    for i in range(width):
        h = raw_headers[i] if i < len(raw_headers) else None
        h = str(h) if h is not None else f"Колонка {get_column_letter(i + 1)}"
        if h in seen:
            seen[h] += 1
            h = f"{h} ({seen[h]})"
        else:
            seen[h] = 1
        headers.append(h)
    data = []
    for offset, r in enumerate(rows[start + 1:], start=start + 2):
        r = (list(r) + [None] * width)[:width]
        if any(c is not None for c in r):
            data.append((offset, r))
        if len(data) > MAX_ROWS:
            raise TableError(f"В листе больше {MAX_ROWS} строк — разбейте файл на части")
    return {"headers": headers, "rows": data}


def read_table(path: Path) -> dict[str, dict]:
    """Возвращает {имя листа: {"headers": [...], "rows": [(номер строки в файле, [значения])]}}."""
    suffix = path.suffix.lower()
    if suffix in (".geojson", ".json"):
        from .geo import GeometryError, read_geojson

        try:
            return read_geojson(path)
        except GeometryError as e:
            raise TableError(str(e)) from e
    sheets: dict[str, list[list]] = {}
    try:
        if suffix in (".xlsx", ".xlsm"):
            wb = load_workbook(path, read_only=True, data_only=True)
            for ws in wb.worksheets:
                sheets[ws.title] = [list(r) for r in ws.iter_rows(values_only=True)]
            wb.close()
        elif suffix == ".xls":
            import xlrd

            book = xlrd.open_workbook(str(path))
            for sh in book.sheets():
                rows = []
                for i in range(sh.nrows):
                    row = []
                    for c in sh.row(i):
                        if c.ctype == xlrd.XL_CELL_DATE:
                            row.append(xlrd.xldate.xldate_as_datetime(c.value, book.datemode))
                        elif c.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                            row.append(None)
                        else:
                            row.append(c.value)
                    rows.append(row)
                sheets[sh.name] = rows
        elif suffix == ".csv":
            raw = path.read_bytes()
            for enc in ("utf-8-sig", "cp1251"):
                try:
                    text = raw.decode(enc)
                    break
                except UnicodeDecodeError:
                    continue
            else:
                raise TableError("Не удалось определить кодировку CSV (ожидается UTF-8 или Windows-1251)")
            sample = text[:20000]
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
            except csv.Error:
                delimiter = ";" if sample.count(";") > sample.count(",") else ","
                dialect = type("Fallback", (csv.excel,), {"delimiter": delimiter})
            sheets["CSV"] = [row for row in csv.reader(io.StringIO(text), dialect)]
        else:
            raise TableError("Поддерживаются файлы .xlsx, .xls, .csv и .geojson")
    except TableError:
        raise
    except Exception as e:  # повреждённый файл, не тот формат и т.п.
        raise TableError(f"Не удалось прочитать файл: {e}") from e
    return {name: _split_header(rows) for name, rows in sheets.items()}


def _autosize(ws, widths=None):
    for i, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 200)), start=1):
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)


def _header_row(ws, headers):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E3EEF9")
    ws.freeze_panes = "A2"


def to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_template(types: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Объекты"
    _header_row(ws, TEMPLATE_HEADERS)
    example_type = types[0]["name"] if types else ""
    ws.append(["OBJ-0001", "Д-2026/015", "77:01:0001001:1234", example_type, "Пример объекта", 55.751244, 37.618423,
               "Москва, ул. Тверская, 1", 500, 12500.5, ""])
    tw = wb.create_sheet("Типы")
    _header_row(tw, ["Код", "Название", "Есть радиус", "Радиус по умолчанию, м"])
    for t in types:
        tw.append([t["code"], t["name"], "да" if t["has_radius"] else "нет", t["default_radius_m"]])
    if types:
        dv = DataValidation(type="list", formula1=f"=Типы!$B$2:$B${len(types) + 1}", allow_blank=False)
        dv.error = "Выберите тип из справочника (лист «Типы»)"
        dv.errorTitle = "Неизвестный тип"
        ws.add_data_validation(dv)
        dv.add("D2:D100000")
    help_ws = wb.create_sheet("Инструкция")
    for line in [
        "Обязательные колонки: Тип, Название, Широта, Долгота.",
        "Координаты — десятичные градусы WGS-84 (например 55,751244). Разделитель — точка или запятая.",
        "ID — внешний идентификатор; нужен для режима «Обновить / добавить».",
        "Номер договора и кадастровый номер — текст; кадастровый номер в формате 77:01:0001001:1234.",
        "Стоимость, тыс. руб. — число не меньше 0, можно дробное (12500,5 = 12,5 млн руб.).",
        "Полигон — необязательная колонка «Геометрия (GeoJSON)» с Polygon/MultiPolygon в WGS-84; тогда координаты можно "
        "не заполнять — маркер встанет посередине. Полигоны удобнее загружать файлом .geojson.",
        "Радиус, м — целое число от 1 до 100 000; учитывается только для типов с зоной. Пусто — радиус типа.",
        "Любые другие колонки сохраняются как дополнительные поля и видны в карточке объекта.",
    ]:
        help_ws.append([line])
    help_ws.column_dimensions["A"].width = 110
    _autosize(ws)
    _autosize(tw)
    return to_bytes(wb)


def _cell(v):
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    return v


GEOMETRY_HEADER = "Геометрия (GeoJSON)"
EXCEL_CELL_LIMIT = 32000


def build_export(objects: list[dict]) -> bytes:
    attr_keys: list[str] = []
    with_geometry = any(o.get("geometry") for o in objects)
    seen = set()
    for o in objects:
        for k in o["attributes"]:
            if k not in seen:
                seen.add(k)
                attr_keys.append(k)
    wb = Workbook(write_only=False)
    ws = wb.active
    ws.title = "Объекты"
    _header_row(ws, TEMPLATE_HEADERS + ([GEOMETRY_HEADER] if with_geometry else []) + attr_keys)
    for o in objects:
        geom = []
        if with_geometry:
            text = json.dumps(o["geometry"], separators=(",", ":")) if o.get("geometry") else None
            # В ячейку Excel помещается ~32 тыс. символов: большие контуры — только в выгрузке GeoJSON.
            geom = [text if not text or len(text) <= EXCEL_CELL_LIMIT else None]
        ws.append(
            [o["external_id"], o["contract_number"], o["cadastral_number"], o["type_name"], o["name"], o["lat"],
             o["lon"], o["address"], o["radius_m"], o["cost"], o["description"]]
            + geom
            + [_cell(o["attributes"].get(k)) for k in attr_keys]
        )
    _autosize(ws)
    if with_geometry:
        ws.column_dimensions[get_column_letter(len(TEMPLATE_HEADERS) + 1)].width = 30
    return to_bytes(wb)


def build_geojson(objects: list[dict]) -> dict:
    """GeoJSON для обмена: свойства названы как колонки шаблона — файл можно загрузить обратно импортом."""
    features = []
    for o in objects:
        props = {
            "ID": o["external_id"], "Номер договора": o["contract_number"], "Кадастровый номер": o["cadastral_number"],
            "Тип": o["type_name"], "Название": o["name"], "Адрес": o["address"], "Радиус, м": o["radius_m"],
            "Стоимость, тыс. руб.": o["cost"], "Описание": o["description"],
        }
        props.update({k: _cell(v) for k, v in o["attributes"].items() if k not in props})
        geometry = o.get("geometry") or {"type": "Point", "coordinates": [o["lon"], o["lat"]]}
        features.append({"type": "Feature", "geometry": geometry, "properties": props})
    return {"type": "FeatureCollection", "features": features}


def build_error_report(headers: list[str], rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ошибки"
    _header_row(ws, ["Строка", "Уровень", "Сообщение"] + headers)
    red = PatternFill("solid", fgColor="FDECEA")
    for r in rows:
        for level, msgs in (("Ошибка", r["errors"]), ("Предупреждение", r["warnings"])):
            for m in msgs:
                ws.append([r["row"], level, m] + [_cell(v) for v in r["values"]])
                if level == "Ошибка":
                    for c in ws[ws.max_row][:3]:
                        c.fill = red
    _autosize(ws)
    return to_bytes(wb)
